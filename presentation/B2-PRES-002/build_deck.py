"""B2-PRES-002 supervisor progress deck. Numbers are copied from the repo files cited in the speaker notes
(no training, no new experiments). Style follows the Slidesgo 'Scientific Research Thesis Defense' template:
16:9, white, navy 0E2A47, grey-blue 869FB2, teal 0097A7, Archivo headings, Manrope body."""
from pathlib import Path
import numpy as np
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).parent
REPO = HERE.parents[1]
NAVY, GREY, TEAL, LIGHT = RGBColor(0x0E, 0x2A, 0x47), RGBColor(0x86, 0x9F, 0xB2), RGBColor(0x00, 0x97, 0xA7), RGBColor(0xEE, 0xF2, 0xF5)
HF, BF = "Archivo", "Manrope"
FIG = HERE / "figures"; FIG.mkdir(exist_ok=True)

# ---------------------------------------------------------------- figures (existing data only)
plt.rcParams.update({"font.size": 11, "axes.spines.top": False, "axes.spines.right": False})
# (a) fitted decay per arm, seed 1234 (B2-E01S_SEED_RESULTS / B2-E02_RESULTS / B2-E03_RESULTS)
arms = [("B1", 8.98), ("B1-FP64", 8.98), ("mixed", 13.24), ("modal (diag.)", 5.40),
        ("full-field\nAdam→\nL-BFGS", 15.64), ("full-field\nfull-batch\nAdam", 12.69), ("modal\nAdam→\nL-BFGS", 3.70)]
fig, ax = plt.subplots(figsize=(8, 3.6))
cols = ["#869FB2"] * 6 + ["#0097A7"]
ax.bar([a for a, _ in arms], [v for _, v in arms], color=cols)
ax.axhline(3.54, color="#0E2A47", ls="--", lw=1.5); ax.set_xlim(-0.6, 7.6)
ax.text(6.45, 3.54 + 0.25, "exact 3.54", ha="left", va="bottom", color="#0E2A47", fontsize=10)
for i, (_, v) in enumerate(arms): ax.text(i, v + 0.3, f"{v:.2f}", ha="center", fontsize=10)
ax.set_ylabel("fitted decay [1/s]"); ax.set_title("Seed 1234, 640k PDE evaluations", fontsize=11, loc="left")
plt.xticks(fontsize=9); fig.tight_layout(); fig.savefig(FIG / "a_decay_per_arm.png", dpi=200); plt.close(fig)
# (b) a-priori excess decay vs window length (docs/hypotheses/B2-E05.md section 1)
Tw, dstar = [1.0, 0.25, 0.10, 0.05], [4.42, 4.04, 2.27, 1.21]
fig, ax = plt.subplots(figsize=(5.2, 3.4))
ax.plot(Tw, dstar, "o-", color="#0097A7", lw=2)
ax.plot([1.0], [5.44], "s", color="#0E2A47", ms=8, label="W1000 measured 5.44")
for x, y in zip(Tw, dstar): ax.annotate(f"{y:.2f}", (x, y), textcoords="offset points", xytext=(6, -12), fontsize=10)
ax.set_xscale("log"); ax.set_xticks(Tw); ax.set_xticklabels(["1.0", "0.25", "0.10", "0.05"])
ax.set_xlabel("window length T_w [s]"); ax.set_ylabel("excess decay d [1/s]"); ax.legend(frameon=False, fontsize=9)
ax.set_title("A-priori illustration (R at B1 level), not a result", fontsize=10, loc="left")
fig.tight_layout(); fig.savefig(FIG / "b_window_law.png", dpi=200); plt.close(fig)
# (c) mode-1 envelope, B1 640k seed 1234 (results_batch2/B2-E04/units/FF_b1-640k_s1234.npz)
z = np.load(REPO / "results_batch2/B2-E04/units/FF_b1-640k_s1234.npz")
fig, ax = plt.subplots(figsize=(5.2, 3.4))
ax.plot(z["A_t"], z["A_a1ex"], color="#869FB2", lw=1.2, label="exact a₁(t)")
ax.plot(z["A_t"], z["A_a"][0], color="#0E2A47", lw=1.2, label="B1 a₁(t)")
ax.set_xlabel("t [s]"); ax.set_ylabel("mode-1 amplitude"); ax.legend(frameon=False, fontsize=9, loc="upper right")
fig.tight_layout(); fig.savefig(FIG / "c_mode1_envelope.png", dpi=200); plt.close(fig)

# ---------------------------------------------------------------- deck helpers
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = prs.slide_width


def text(slide, x, y, w, h, runs, size=16, color=NAVY, font=BF, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, r in enumerate(runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = r; p.space_after = Pt(6)
        for run in p.runs:
            run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
            run.font.color.rgb = color
    return tb


def slide(title, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    if kicker:
        text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), 12, TEAL, HF, True)
    text(s, 0.6, 0.65, 12.2, 1.2, title, 28, NAVY, HF, True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def bullets(s, items, x=0.6, y=2.0, w=12, h=4.8, size=18):
    return text(s, x, y, w, h, ["•  " + i for i in items], size)


def table(s, rows, x=0.6, y=2.0, w=12.1, colw=None, size=13):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows))).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v)
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else RGBColor(255, 255, 255))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.name = Pt(size), (HF if i == 0 else BF)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else NAVY
    if colw:
        for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    return t


def tag(s, label, x=10.6, y=0.3):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.4), Inches(0.42))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xED, 0x56, 0x1B); b.line.fill.background()
    tf = b.text_frame; tf.text = label
    for r in tf.paragraphs[0].runs: r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(12), True, HF, RGBColor(255, 255, 255)


# ---------------------------------------------------------------- slides
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.55), Inches(1.6), Inches(0.08)); acc.fill.solid(); acc.fill.fore_color.rgb = TEAL; acc.line.fill.background()
text(s, 0.8, 1.6, 11.5, 2.8, "PINN optimisation for damped structural vibration: findings and path to railway predictive maintenance", 38, RGBColor(255, 255, 255), HF, True)
text(s, 0.8, 4.85, 10, 0.6, "Shreyansh  ·  supervisor progress review", 20, GREY, BF)

s = slide("One objective, three papers: understand PINN failure on a controlled beam, then transfer to railways", "Objective")
table(s, [["Paper", "Scope", "Status"],
          ["P1  PINN optimisation (beam)", "Why long-window damped-vibration PINNs fail; window law + reference-free rule", "Diagnostics done (E01–E04); window sweep paused"],
          ["P2  Railway transfer + Physics Reference Layer", "Verification-labelled physics references for railway models", "Not started"],
          ["P3  Railway application", "Predictive maintenance use case (axle / bearing / wheel)", "Not started"]], y=2.1, colw=[3.6, 5.6, 2.9], size=15)
text(s, 0.6, 4.6, 12, 1, "No railway claim is made yet: all evidence so far is on the beam benchmark.", 16, GREY)

s = slide("Benchmark: a damped fixed-fixed Euler–Bernoulli beam with an exact analytical solution", "Benchmark FE-D-M1",
          "Sources: B2-E01_RESULTS.csv (omega_d 129.32 rad/s, full_window_cycles 20.58); B2-E05.md (zeta = gamma/2 = 3.54).")
table(s, [["Quantity", "Value"], ["Boundary conditions", "fixed-fixed, mode 1"], ["Damped frequency ω_d", "129.32 rad/s (20.58 Hz)"],
          ["Window", "1 s = 20.58 cycles"], ["Exact decay rate ζ", "3.54 1/s"], ["Reference", "analytical modal solution (exact)"],
          ["Success = ", "full-window persistence (20.58 cycles) + correct decay + low L2"]], y=2.0, w=8.5, colw=[3.2, 5.3], size=15)

s = slide("Batch 1's best configuration (B1) reproduces the frequency but over-damps and collapses at 8.1 of 20.6 cycles", "Batch 1 summary",
          "Source: configs/frozen/batch1_hard_tanh2.yaml batch1_result. The number of Batch-1 runs is not stated in the listed sources and is left out.")
table(s, [["B1 = hard IC/BC + tanh²(ω₁t) + D4 Fourier + exp-decay LR + mini-batch 128 (Z4-20K)", "Value"],
          ["Training", "20,000 steps, 2.56M PDE evaluations, 2,290 s"], ["L2 (exact)", "0.263"], ["Fitted frequency", "129.2 rad/s"],
          ["Collapse", "0.393 s = 8.1 of 20.6 cycles"], ["Decay", "5.40 vs 3.54 1/s (over-damped)"], ["Verdict", "FAIL: no validated optimised PINN"]],
      y=2.0, colw=[8.0, 4.1], size=15)

s = slide("Five candidate explanations were tested at matched compute and eliminated", "Hypotheses eliminated",
          "Sources: B2-E01_SEED_REPLICATION.md, B2-E02_REPORT.md, B2-E03_REPORT.md, B2-E03_RESULTS.csv. Network capacity: no Batch-2 test in the listed sources, so it is not listed.")
table(s, [["Hypothesis", "Test", "Key number", "Verdict"],
          ["Floating-point precision", "FP64 control (E01)", "identical L2 0.526 / 0.549 / 0.402 in 3 seeds", "not material"],
          ["4th-order operator", "mixed form, order 2 (E01)", "persistence 1.67–2.13 vs B1 3.11–5.10 cycles", "cheaper, not better"],
          ["Collocation coverage", "resampling / 4× reservoir (E03)", "H1 supported in 1/3 seeds only; 4X leaves front stalled", "insufficient (CASE D)"],
          ["Optimiser: full-batch Adam", "ADAM-FULL (E03)", "H2 falsified 0/3", "eliminated"],
          ["L-BFGS on the full field", "Adam→L-BFGS transfer (E02)", "worse than B1 in 3/3; L2 ×1.28–1.54", "eliminated"]],
      y=2.0, colw=[2.7, 3.0, 4.4, 2.0], size=13)

s = slide("The same optimiser solves the modal problem but fails on the full field", "Optimiser reversal",
          "Sources: B2-E02_RESULTS.csv, B2-E02T_RESULTS.csv, B2-E02_REPORT.md.")
table(s, [["Seed", "Modal Adam→L-BFGS: cycles / L2", "Full-field B1: cycles / L2", "Full-field Adam→L-BFGS: cycles / L2"],
          ["1234", "20.58 / 0.038", "3.15 / 0.526", "1.60 / 0.671"], ["1235", "20.58 / 0.071", "3.11 / 0.549", "0.72 / 0.843"],
          ["1236", "20.58 / 0.093", "5.10 / 0.402", "2.53 / 0.578"]], y=2.0, colw=[1.4, 3.6, 3.4, 3.7], size=15)
bullets(s, ["Modal: full window (20.58 cycles) in 3/3 seeds — but the exact mode shape is supplied (diagnostic, not a method).",
            "Full field: worse than B1 in 3/3 seeds (persistence −1.6 to −2.6 cycles).",
            "Decision E: the improvement exists only in the modal diagnostic."], y=4.0, size=16)

s = slide("Collapse is a loss of first-mode amplitude; no candidate diagnostic leads it (CASE F)", "B2-E04 diagnostics",
          "Sources: B2-E04_REPORT.md; B2-E04/units/FF_b1-640k_s1234.json (rho_nonfund 1.14e-6) and .npz (A_t, A_a, A_a1ex).")
s.shapes.add_picture(str(FIG / "c_mode1_envelope.png"), Inches(7.3), Inches(1.9), width=Inches(5.6))
bullets(s, ["Residual energy outside mode 1: 1.1·10⁻⁶ (B1, 640k, seed 1234); higher modes ≤ 5·10⁻⁵ of exact energy in all runs.",
            "Collapse = a₁(t) → 0, not growth of higher modes.",
            "Spatial/temporal gradient ratio: 1.6–3.4·10⁴ at init, 0.87–2.0 after training (no late imbalance).",
            "Verdict: CASE F — no clear mechanism among A–D (B partial, not discriminating)."], w=6.5, size=15)

s = slide("Every failing arm decays too fast: artificial damping is the common symptom", "Mechanism",
          "Sources: B2-E01S_SEED_RESULTS / B2-E01_SEED_REPLICATION.md, B2-E02_RESULTS.csv, B2-E03_RESULTS.csv (fit_decay, seed 1234). Law: sprint/SPRINT_24H_v2_CLAUDE_CODE.md section 0 (proposed, not validated).")
s.shapes.add_picture(str(FIG / "a_decay_per_arm.png"), Inches(0.5), Inches(1.9), width=Inches(7.6))
text(s, 8.4, 2.0, 4.5, 4.5, ["Proposed damping-bias law (long windows):", "λ* = √(ζ² + (R / 2ω)²)",
                             "R = unresolved residual per unit amplitude [s⁻²], measurable without the reference.",
                             "Only the modal L-BFGS arm (3.70) is near the exact 3.54.",
                             "Status: hypothesis under test (B2-E05)."], 16)

s = slide("Reweighting the loss moves the damping in both directions, but none beat B1", "Exploratory — seed 1234 only",
          "Sources: sprint/SPRINT_24H_v2_CLAUDE_CODE.md section 1; sprint/exploratory_runs_2026-10-09.csv. Scratch runs, not results of record.")
tag(s, "EXPLORATORY")
table(s, [["Variant", "L2_exact", "Fitted decay (exact 3.54)", "Behaviour"],
          ["B1 (logged)", "0.526", "8.98", "over-damped"],
          ["normalised loss (degree 0)", "5.48 @2500 steps", "growing", "amplitude ratio 1.8 / 9.3 / 61 at t = 0.05 / 0.4 / 1.0 s"],
          ["fixed exponent p = 0.5", "0.818", "10.20", "late spurious growth"],
          ["energy-rate loss μ = 10", "4.49 @2500 steps", "—", "collapsed to a static field"],
          ["feedback controller on p", "0.604", "11.19", "sensor latched onto growth; p driven to 0"]],
      y=2.0, colw=[3.2, 2.2, 2.6, 4.1], size=14)
text(s, 0.6, 5.0, 12, 1, "Takeaway: the loss's amplitude homogeneity sets the sign of the bias; the remaining lever with a quantitative prediction is window length.", 15, GREY)

s = slide("Window law: shorter windows should cut excess decay; first run W1000 reproduces B1 bit-for-bit", "Window law and prior art",
          "Sources: docs/hypotheses/B2-E05.md section 1 (a-priori table); results_batch2/runs/B2-E05-W1000-s1234/results_row.csv; sprint/STATUS.md (W1000 predicted 6.29).")
s.shapes.add_picture(str(FIG / "b_window_law.png"), Inches(0.5), Inches(1.9), width=Inches(6.2))
text(s, 7.0, 2.0, 6.0, 4.8, ["W1000 (1 window, seed 1234): bit-identical to logged B1, L2 0.52639.",
                             "Measured excess decay 5.44 1/s vs predicted 6.29 from its measured R (2374 s⁻²).",
                             "Prior art: AT-PINN (2024) and AT-PINN-HC (2025) do time-marching with hard constraints on EB beams, but give no relation between segment length and damping error.",
                             "Novelty claimed only for the law + reference-free window rule, if validated."], 15)

s = slide("Status: diagnostics are validated; the window law is untested beyond one window", "Honest status")
table(s, [["Validated (pre-registered, 3 seeds)", "Exploratory (seed 1234)", "Pending"],
          ["B1 reproduced bit-exactly in PINN-RP", "loss-homogeneity runs (6 variants)", "W050 / W100 / W250 sweep: not yet run (W050 interrupted at 11/20 windows)"],
          ["FP64, mixed, coverage, L-BFGS transfer eliminated", "W1000 vs prediction (1 run)", "Block-3 gate: prediction vs measurement"],
          ["Collapse = mode-1 amplitude loss (E04, CASE F)", "", "Seeds 1235/1236; sweep currently PAUSED awaiting PI"]],
      y=2.0, colw=[4.2, 3.6, 4.3], size=14)
text(s, 0.6, 5.3, 12, 1, "No optimised PINN is claimed.", 16, RGBColor(0xED, 0x56, 0x1B), HF, True)

s = slide("Timeline: P1 manuscript by 20 Oct, P2 and P3 by 20 Nov, each behind a gate", "Timeline")
table(s, [["Milestone", "Date", "Gate"],
          ["Window sweep (W050/W100/W250) + Block-3 analysis", "on resume", "PASS / PARTIAL / FAIL pre-registered in B2-E05.md"],
          ["P1 manuscript (beam)", "20 Oct", "PASS → law + rule; FAIL → diagnostic manuscript"],
          ["P2 railway transfer + Physics Reference Layer", "by 20 Nov", "supervisor decisions (next slide) + data access"],
          ["P3 railway application", "by 20 Nov", "validated transfer model"]], y=2.0, colw=[5.2, 1.9, 5.0], size=15)

s = slide("Railway transfer model: axle, bearing and wheel dynamics", "Placeholder")
text(s, 0.6, 2.4, 12, 1, "to be added", 24, GREY)

s = slide("Decisions needed from the supervisor", "Next steps")
bullets(s, ["Target vehicle", "Sensor scenario: on-board or wayside", "Bearing type and drawing", "Validation data source",
            "Does a “submitted” paper count for graduation?"], size=22)

out = HERE / "PINN_progress.pptx"; prs.save(out); print(out)
